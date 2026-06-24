import pandas as pd
from datetime import datetime
from pathlib import Path
import os

class ExcelExporter:
    """Clase para exportar datos estadísticos a archivos Excel"""
    
    def __init__(self, output_dir="exports"):
        """
        Inicializa el exportador de Excel
        
        Args:
            output_dir: Directorio donde se guardarán los archivos Excel
        """
        self.output_dir = output_dir
        self._create_output_dir()
    
    def _create_output_dir(self):
        """Crea el directorio de salida si no existe"""
        Path(self.output_dir).mkdir(parents=True, exist_ok=True)
    
    def export_values_to_excel(self, values_data, filename=None):
        """
        Exporta los valores de la base de datos a un archivo Excel
        
        Args:
            values_data: Lista de tuplas con los datos (id, value, session_id, description, created_at)
            filename: Nombre del archivo (opcional, se genera automáticamente si no se proporciona)
        
        Returns:
            str: Ruta completa del archivo generado
        """
        if not values_data:
            print("No hay datos para exportar.")
            return None
        
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"valores_estadisticos_{timestamp}.xlsx"
        
        filepath = os.path.join(self.output_dir, filename)
        
        # Crear DataFrame con los datos
        df = pd.DataFrame(
            values_data,
            columns=["ID", "Valor", "Session ID", "Descripción", "Fecha de Creación"]
        )
        
        # Exportar a Excel con formato mejorado
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Datos')
            
            # Obtener el workbook y la hoja
            workbook = writer.book
            worksheet = writer.sheets['Datos']
            
            # Ajustar el ancho de las columnas
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                worksheet.column_dimensions[column_letter].width = adjusted_width
            
            # Dar formato a la fila de encabezados
            from openpyxl.styles import Font, PatternFill, Alignment
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
        
        print(f"✓ Datos exportados exitosamente a: {filepath}")
        return filepath
    
    def export_statistics_summary(self, stats_data, filename=None):
        """
        Exporta un resumen estadístico a un archivo Excel
        
        Args:
            stats_data: Diccionario con las estadísticas (average, variance, deviation, count)
            filename: Nombre del archivo (opcional)
        
        Returns:
            str: Ruta completa del archivo generado
        """
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"resumen_estadistico_{timestamp}.xlsx"
        
        filepath = os.path.join(self.output_dir, filename)
        
        # Crear datos del resumen
        summary_data = {
            'Estadística': list(stats_data.keys()),
            'Valor': list(stats_data.values())
        }
        
        df = pd.DataFrame(summary_data)
        
        # Exportar a Excel con formato
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Resumen')
            
            workbook = writer.book
            worksheet = writer.sheets['Resumen']
            
            # Ajustar ancho de columnas
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                worksheet.column_dimensions[column_letter].width = adjusted_width
            
            # Dar formato a encabezados
            from openpyxl.styles import Font, PatternFill, Alignment
            header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
        
        print(f"✓ Resumen estadístico exportado a: {filepath}")
        return filepath
    
    def export_combined_report(self, values_data, stats_data, filename=None):
        """
        Exporta un reporte combinado con datos y estadísticas en hojas diferentes
        
        Args:
            values_data: Lista de tuplas con los valores
            stats_data: Diccionario con las estadísticas
            filename: Nombre del archivo (opcional)
        
        Returns:
            str: Ruta completa del archivo generado
        """
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"reporte_completo_{timestamp}.xlsx"
        
        filepath = os.path.join(self.output_dir, filename)
        
        # Crear DataFrames
        df_valores = pd.DataFrame(
            values_data,
            columns=["ID", "Valor", "Session ID", "Descripción", "Fecha de Creación"]
        )
        
        df_stats = pd.DataFrame({
            'Estadística': list(stats_data.keys()),
            'Valor': list(stats_data.values())
        })
        
        # Exportar a Excel con múltiples hojas
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            df_valores.to_excel(writer, index=False, sheet_name='Datos')
            df_stats.to_excel(writer, index=False, sheet_name='Estadísticas')
            
            workbook = writer.book
            
            # Formatear hoja de Datos
            worksheet_datos = writer.sheets['Datos']
            self._format_worksheet(worksheet_datos, "4472C4")
            
            # Formatear hoja de Estadísticas
            worksheet_stats = writer.sheets['Estadísticas']
            self._format_worksheet(worksheet_stats, "70AD47")
        
        print(f"✓ Reporte completo exportado a: {filepath}")
        return filepath
    
    def _format_worksheet(self, worksheet, header_color):
        """Aplica formato a una hoja de trabajo"""
        from openpyxl.styles import Font, PatternFill, Alignment
        
        # Ajustar ancho de columnas
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            worksheet.column_dimensions[column_letter].width = adjusted_width
        
        # Dar formato a encabezados
        header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
