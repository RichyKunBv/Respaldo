import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

def main():
    input_file = "datos_simulacion.csv"
    
    print(f"Leyendo {input_file}...")
    
    try:
        df = pd.read_csv(input_file)
    except FileNotFoundError:
        print(f"No esta '{input_file}'")
        return
    except Exception as e:
        print(f"Error de datos: {e}")
        return

    output_file = 'Resultados.xlsx'
    
    print("Exportando a Excel...")
    df.to_excel(output_file, index=False, sheet_name="Hoja1")
    
    print("Metiendole la mano a excel...")
    wb = load_workbook(output_file)
    ws = wb["Hoja1"]
    
    total_datos = len(df)
    last_row = total_datos + 1
    prom_row = last_row + 2
    desv_row = last_row + 3
    ic_sup_row = last_row + 5
    ic_inf_row = last_row + 6
    
    ws[f'C{prom_row}'] = 'Promedio de Tráfico:'
    ws[f'D{prom_row}'] = f'=AVERAGE(A2:A{last_row})'
    
    ws[f'C{desv_row}'] = 'Desviacion (Sigma):'
    ws[f'D{desv_row}'] = f'=STDEV.S(A2:A{last_row})'

    ws[f'C{ic_sup_row}'] = 'Limite Sup IC (95%):'
    ws[f'D{ic_sup_row}'] = f'=D{prom_row} + 1.96*(D{desv_row}/SQRT({total_datos}))'
    
    ws[f'C{ic_inf_row}'] = 'Limite Inf IC (95%):'
    ws[f'D{ic_inf_row}'] = f'=D{prom_row} - 1.96*(D{desv_row}/SQRT({total_datos}))'
    
    print("Gaus")

    columna_datos = df.columns[0]
    datos_limpios = pd.to_numeric(df[columna_datos], errors='coerce').dropna()
    
    counts, bins = np.histogram(datos_limpios, bins=50)
    
    ws['J1'] = 'Rango de Tráfico'
    ws['K1'] = 'Frecuencia (Campana)'
    
    for i in range(len(counts)):
        punto_medio = 0.5 * (bins[i] + bins[i+1])
        ws.cell(row=i+2, column=10, value=punto_medio)
        ws.cell(row=i+2, column=11, value=counts[i])
    
    print("grafiko")
    chart = LineChart()
    chart.title = "Distribución del Tráfico (Gauss)"
    chart.style = 27  
    chart.width = 18
    chart.height = 10
    
    data = Reference(ws, min_col=11, min_row=1, max_row=len(counts)+1)
    chart.add_data(data, titles_from_data=True)
    
    ws.add_chart(chart, "E10")
    
    wb.save(output_file)
    print(f"¡Terminado! Checa tu {output_file}")

if __name__ == "__main__":
    main()