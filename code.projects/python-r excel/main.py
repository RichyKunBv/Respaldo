import pandas as pd
import rpy2.robjects as robjects
from rpy2.robjects import pandas2ri
from rpy2.robjects.conversion import localconverter
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

def main():
    print("Conectando con R y cargando el script 'simulacion.R'...")
    robjects.r.source("simulacion.R")
    
    print("Generando 100,000 datos de simulación...")
    generar_datos = robjects.globalenv['generar_datos']
    r_dataframe = generar_datos(100000)
    
    with localconverter(robjects.default_converter + pandas2ri.converter):
        df = robjects.conversion.rpy2py(r_dataframe)
        
    output_file = 'Resultados_Simulacion_100k.xlsx'
    
    # 1. Guardar a excel
    print("Exportando a Excel...")
    df.to_excel(output_file, index=False, sheet_name="Hoja1")
    
    # 2. Agregar cosas visuales arre
    print("Agregando fórmulas, parámetros y el gráfico...")
    wb = load_workbook(output_file)
    ws = wb["Hoja1"]
    
    # Datos del excel origina
    ws['F1'] = 'Demanda media'
    ws['G1'] = 10000
    ws['F2'] = 'Desviacion'
    ws['G2'] = 2000
    ws['F4'] = 'Precio Venta'
    ws['G4'] = 50
    ws['H4'] = 70
    ws['F6'] = 'Costo Unitario'
    ws['G6'] = 30
    ws['F7'] = 'Desviacion'
    ws['G7'] = 5
    
    # cosas de hasta abajo
    last_row = 100001
    prom_row = last_row + 2
    desv_row = last_row + 3
    
    ws[f'A{prom_row}'] = f'=AVERAGE(A2:A{last_row})'
    ws[f'A{desv_row}'] = f'=STDEV.S(A2:A{last_row})'
    
    ws[f'C{prom_row}'] = f'=AVERAGE(C2:C{last_row})'
    ws[f'C{desv_row}'] = f'=STDEV.S(C2:C{last_row})'
    
    # grafico
    chart = LineChart()
    chart.title = "Título del gráfico"
    chart.style = 27  
    chart.width = 18
    chart.height = 10
    
    # Demanda
    data = Reference(ws, min_col=1, min_row=1, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    
    # Insertar gráfico
    ws.add_chart(chart, "E10")
    
    wb.save(output_file)
    print(f"¡Listo! Revisa tu archivo {output_file}")

if __name__ == "__main__":
    main()
