import pandas as pd
import numpy as np
import rpy2.robjects as robjects
from rpy2.robjects import pandas2ri
from rpy2.robjects.conversion import localconverter
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

def main():
    print("agarrando a R...")
    robjects.r.source("simulacion.R")
    
    print("generando los 100k datos...")
    generar_datos = robjects.globalenv['generar_datos']
    r_dataframe = generar_datos(100000)
    
    with localconverter(robjects.default_converter + pandas2ri.converter):
        df = robjects.conversion.rpy2py(r_dataframe)
        
    output_file = 'Resultados_Simulacion_100k.xlsx'
    
    print("exportando a excel...")
    df.to_excel(output_file, index=False, sheet_name="Hoja1")
    
    print("metiendole mano al excel...")
    wb = load_workbook(output_file)
    ws = wb["Hoja1"]
    
    ws['F1'] = 'Demanda media'
    ws['G1'] = 10000
    ws['F2'] = 'Desviacion'
    ws['G2'] = 2000
    ws['F4'] = 'Precio Venta'
    ws['G4'] = 50
    ws['H4'] = 70
    ws['F6'] = 'Costo Unitario'
    ws['G6'] = 30
    ws['F7'] = 'Desviacion'
    ws['G7'] = 5
    
    last_row = 100001
    prom_row = last_row + 2
    desv_row = last_row + 3
    ic_sup_row = last_row + 5
    ic_inf_row = last_row + 6
    
    ws[f'A{prom_row}'] = f'=AVERAGE(A2:A{last_row})'
    ws[f'A{desv_row}'] = f'=STDEV(A2:A{last_row})'
    ws[f'C{prom_row}'] = f'=AVERAGE(C2:C{last_row})'
    ws[f'C{desv_row}'] = f'=STDEV(C2:C{last_row})'

    ws[f'C{prom_row}'] = 'Promedio Ganancia:'
    ws[f'D{prom_row}'] = f'=AVERAGE(D2:D{last_row})'
    
    ws[f'C{desv_row}'] = 'Desviacion Ganancia:'
    ws[f'D{desv_row}'] = f'=STDEV(D2:D{last_row})'

    # 95
    ws[f'C{ic_sup_row}'] = 'Limite Sup IC (95%):'
    ws[f'D{ic_sup_row}'] = f'=D{prom_row} + 1.96*(D{desv_row}/SQRT(100000))'
    
    ws[f'C{ic_inf_row}'] = 'Limite Inf IC (95%):'
    ws[f'D{ic_inf_row}'] = f'=D{prom_row} - 1.96*(D{desv_row}/SQRT(100000))'
    

    counts, bins = np.histogram(df['Demanda'], bins=50)
    
    ws['J1'] = 'Rango Demanda'
    ws['K1'] = 'Frecuencia (Campana)'
    
    for i in range(len(counts)):
        punto_medio = 0.5 * (bins[i] + bins[i+1])
        ws.cell(row=i+2, column=10, value=punto_medio)
        ws.cell(row=i+2, column=11, value=counts[i])
    
    # grafico mamalon :VVV
    chart = LineChart()
    chart.title = "Distribución de la Demanda (Gauss)"
    chart.style = 27  
    chart.width = 18
    chart.height = 10
    
    data = Reference(ws, min_col=11, min_row=1, max_row=len(counts)+1)
    chart.add_data(data, titles_from_data=True)
    
    ws.add_chart(chart, "E10")
    
    wb.save(output_file)
    print(f"{output_file}")

if __name__ == "__main__":
    main()