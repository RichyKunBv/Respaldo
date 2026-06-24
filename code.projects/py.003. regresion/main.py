import oracledb
import numpy as np
from sklearn.linear_model import LinearRegression
from sklearn.metrics import r2_score
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.utils import get_column_letter

# ---------- Conexión a Oracle (tus credenciales) ----------
username = "webo"
password = "escamilla13XD"
host = "192.168.215.2"
port = 1521
service_name = "FREEPDB1"
dsn = oracledb.makedsn(host, port, service_name=service_name)

try:
    conn = oracledb.connect(user=username, password=password, dsn=dsn)
    print("✅ Conectado como WEBO")
except Exception as e:
    print("❌ Error de conexión:", e)
    exit()

# Generar los números en Oracle (sin tabla)
sql_generador = """
    SELECT n,
           (MOD(TRUNC(ABS(SIN(n * 123456.789) * 1000000)), 900000) + 100000) / 1000000 AS valor
    FROM (
        SELECT LEVEL AS n
        FROM DUAL
        CONNECT BY LEVEL <= 10000
    )
    ORDER BY n
"""

cursor = conn.cursor()
cursor.execute(sql_generador)
rows = cursor.fetchall()
cursor.close()
conn.close()

print(f"Se recuperaron {len(rows)} registros desde Oracle.")

# Convertir a arrays
x = np.array([r[0] for r in rows], dtype=float).reshape(-1, 1)
y = np.array([r[1] for r in rows], dtype=float)

# ---------- Regresión lineal ----------
modelo = LinearRegression()
modelo.fit(x, y)

pendiente = modelo.coef_[0]
interseccion = modelo.intercept_
y_pred = modelo.predict(x)
r2 = r2_score(y, y_pred)

print(f"Pendiente: {pendiente:.10f}  Intersección: {interseccion:.10f}  R²: {r2:.6f}")

# ---------- Crear Excel con gráfica ----------
wb = Workbook()
ws = wb.active
ws.title = "Datos y Regresión"

# Escribir encabezados
ws.append(["n", "valor", "prediccion"])

# Escribir datos (n, valor, predicción)
for i, (n, val) in enumerate(rows, start=2):   # fila 2 en adelante
    ws.cell(row=i, column=1, value=n)
    ws.cell(row=i, column=2, value=val)
    ws.cell(row=i, column=3, value=float(y_pred[i-2]))   # y_pred tiene la misma longitud

# Crear gráfico de dispersión
chart = ScatterChart()
chart.title = "Regresión lineal sobre números pseudoaleatorios"
chart.x_axis.title = "n"
chart.y_axis.title = "valor"
chart.style = 13   # estilo visual

# Serie 1: datos reales (nube de puntos)
x_values = Reference(ws, min_col=1, min_row=2, max_row=len(rows)+1)
y_values = Reference(ws, min_col=2, min_row=2, max_row=len(rows)+1)
series_data = Series(y_values, x_values, title="Datos originales")
series_data.marker.symbol = "circle"
series_data.marker.size = 2          # puntos pequeños
series_data.graphicalProperties.line.noFill = True   # sin línea entre puntos
chart.series.append(series_data)

# Serie 2: recta de regresión (usando columna de predicciones)
x_pred = Reference(ws, min_col=1, min_row=2, max_row=len(rows)+1)
y_pred_range = Reference(ws, min_col=3, min_row=2, max_row=len(rows)+1)
series_line = Series(y_pred_range, x_pred, title="Recta de regresión")
series_line.marker.symbol = "none"   # sin marcadores, solo línea
series_line.graphicalProperties.line.width = 20000  # línea gruesa (en EMUs)
chart.series.append(series_line)

# Ajustar colores (opcional)
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.line import LineProperties, LineEndProperties
series_line.graphicalProperties.line.solidFill = "FF0000"  # rojo

# Insertar gráfico en la hoja
ws.add_chart(chart, "E2")

# Ajustar ancho de columnas para que no se solape
ws.column_dimensions[get_column_letter(1)].width = 10
ws.column_dimensions[get_column_letter(2)].width = 12
ws.column_dimensions[get_column_letter(3)].width = 12

# Guardar archivo Excel
wb.save("regresion_grafica.xlsx")
print("\nArchivo 'regresion_grafica.xlsx' guardado con la gráfica incluida.")